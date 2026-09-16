"""Build staff documents with the approved bundled Python runtime.

Screenshots are included unchanged. Render the DOCX with the bundled
Documents/render_docx.py --emit_pdf and inspect each page before distribution.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/cashier'
SHOTS=OUT/'screenshots'
PAGES=[
 ('Choose and review the items','01-choose-items.png',[
  ('1  Sign in','Use your own account. Open B Active cashier, then Open tablet checkout. Check your name. Keep the tablet online.'),
  ('2  Add exact items','Search the product or SKU. Check size, color and quantity against the items in your hand. Tap Add item. Use + or − in the basket to adjust quantity.'),
  ('3  Review the sale','Enter an email only if the customer wants a confirmation. Tap Review sale & payment, then read the final total to the customer.')
 ],'Missing item, incorrect stock or wrong price? Stop and ask your manager.'),
 ('Take cash','03-cash-panel.png',[
  ('1  Count the money','Enter the cash you actually counted under Amount received (PHP). Check the change shown.'),
  ('2  Confirm once','Tap Confirm cash received once. Wait for Payment received on the cashier screen.'),
  ('3  Return the change','Return the displayed change. Continue to the invoice and handover steps on page 4.')
 ],'An error does not mean the sale failed. Refresh the same sale before collecting again.'),
 ('Take a digital payment','05-digital-panel.png',[
  ('1  Open the payment page','Tap Pay digitally. Ask the customer to scan the code with their phone camera. They choose QRPh, Maya, ShopeePay or GrabPay on the secure payment page.'),
  ('2  Wait for confirmation','Keep this screen open. It checks automatically; you can also tap Refresh payment status. Wait for Payment received before continuing to page 4.'),
  ('3  Keep the same sale','If payment remains pending, keep the goods and ask a manager. Do not switch to cash or create a replacement sale.')
 ],'The displayed code opens a website. It is not the bank-app QRPh code. Never use a customer screenshot as proof of payment.'),
 ('Write the invoice and hand over','04-invoice-panel.png',[
  ('1  Check payment','The cashier screen must say Payment received. Recheck the items, sizes, colors and any change.'),
  ('2  Write the invoice','Issue a registered handwritten invoice for every sale. Enter its serial number under Handwritten invoice number.'),
  ('3  Give the customer their purchase','Give the invoice, goods and any change. Tap Confirm invoice & hand over goods as you hand them over.')
 ],'Only the manager may resolve payment or stock exceptions. Never bypass a pending or review message.'),
 ('Finish the sale','06-complete-panel.png',[
  ('1  Check completion','Confirm Sale complete and the invoice number, then tap Start next sale.'),
  ('2  Handle email requests','Submitted to email service means accepted for sending, not delivered. Use Resend email confirmation if needed. Resending does not collect payment again.'),
  ('3  Sign out','Use Sign out when finished. An unfinished sale remains linked to your account; sign back in with the same account to resume it.')
 ],'The email is an order/payment confirmation, not a BIR tax invoice. Always give the handwritten invoice.')
]
for _,name,_,_ in PAGES:
 if not (SHOTS/name).is_file(): raise SystemExit('Missing required screenshot: '+name)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.55);sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name]; st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(20)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
footer=sec.footer.paragraphs[0];footer.text='B Active staff training • Local training examples • '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.style=doc.styles['Normal'];footer.runs[0].font.size=Pt(9)
for index,(title,image,steps,note) in enumerate(PAGES):
 if index:doc.add_page_break()
 if index==0:
  doc.add_paragraph('B Active in store checkout guide','Title')
  doc.add_paragraph('Use after manager release approval. Screens show the implemented local training system with synthetic sales; they do not establish live payment readiness.').runs[0].font.size=Pt(9)
 doc.add_paragraph(title,'Heading 1')
 if index==0:
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(SHOTS/image),width=Inches(6.95))
  p=doc.add_paragraph('Training example — synthetic items and payments.');p.runs[0].italic=True;p.runs[0].font.size=Pt(9)
  for heading,body in steps:
   doc.add_paragraph(heading,'Heading 2');doc.add_paragraph(body)
  p=doc.add_paragraph();r=p.add_run(note);r.bold=True;r.font.size=Pt(10)
 else:
  doc.add_paragraph('Training example — synthetic sale on the implemented cashier screen.').runs[0].font.size=Pt(9)
  table=doc.add_table(rows=1,cols=2);table.autofit=False
  table.columns[0].width=Inches(3.3);table.columns[1].width=Inches(3.6)
  left,right=table.rows[0].cells
  left.width=Inches(3.3);right.width=Inches(3.6)
  for heading,body in steps:
   left.add_paragraph(heading,'Heading 2');left.add_paragraph(body)
  p=left.add_paragraph();r=p.add_run(note);r.bold=True;r.font.size=Pt(10)
  right.paragraphs[0].add_run().add_picture(str(SHOTS/image),width=Inches(3.3))
doc.add_page_break();doc.add_paragraph('Dos and donts','Heading 1')
for title,body in [('Do','Use your own account. Check exact variations. Count cash. Read the final total. Wait for the cashier’s payment confirmation. Issue a registered handwritten invoice for every sale.'),('Do not','Accept payment screenshots. Hand over while pending. Change a pending digital sale to cash. Create a replacement sale after an error. Use a separate QR or payment link. Share accounts or adjust inventory yourself.')]:
 doc.add_paragraph(title,'Heading 2');doc.add_paragraph(body)
doc.add_paragraph('When to stop and ask for help','Heading 1')
rows=[('Waiting for payment','Keep the goods. Refresh this sale. Ask a manager if unresolved.'),('Manager help needed','Stop. Quote the order number. Do not collect again or hand over goods.'),('Connection lost','Keep the page open. Reconnect and refresh the same sale.'),('Check the previous request','Tap Resume this sale. Retry this same sale, if shown, repeats the original reference and items without taking payment.'),('Sale cannot be found','Ask a manager to check your active sale. Do not create a replacement.'),('Wrong basket after review','Use Cancel this unpaid sale only if available. Wait for cancellation before starting the corrected sale. Otherwise ask a manager.'),('Refund exchange or discount','Ask a manager. These actions are not available at the cashier.'),('Email sending failed','Resend the email confirmation. Do not repeat checkout.')]
for heading,body in rows:
 doc.add_paragraph(heading,'Heading 2');doc.add_paragraph(body)
doc.add_paragraph('Always quote the order number when asking your manager for help.').runs[0].bold=True
# Remove inherited template paragraph borders, including the stock blue Title rule.
for tree in (doc._element, doc.styles.element):
 for border in list(tree.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
doc.save(OUT/'cashier-staff-guide.docx')

styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='CounterTitle',fontName='Helvetica-Bold',fontSize=22,leading=25,textColor=colors.black,spaceAfter=12))
styles.add(ParagraphStyle(name='CounterHeading',fontName='Helvetica-Bold',fontSize=12,leading=15,spaceBefore=12,spaceAfter=5))
styles.add(ParagraphStyle(name='CounterBody',fontName='Helvetica',fontSize=11,leading=15,spaceAfter=7))
styles.add(ParagraphStyle(name='CounterNote',fontName='Helvetica',fontSize=9,leading=12,spaceAfter=10,textColor=colors.HexColor('#444444')))
story=[Paragraph('B Active checkout checklist',styles['CounterTitle']),Paragraph('Use after manager release approval. Keep beside the store tablet.',styles['CounterNote'])]
for heading,body in [('1  Check the items','Correct product, size, color and quantity.'),('2  Review the final total','Tap <b>Review sale &amp; payment</b>. Read the final amount to the customer.'),('3  Collect once','<b>Cash:</b> count it, enter amount, confirm. <b>Digital:</b> customer scans with phone camera and pays on the secure page.'),('4  Wait for Payment received','Keep the goods until this cashier screen confirms payment.'),('5  Write the invoice','Issue a registered handwritten invoice for <b>every sale</b>. Enter its number.'),('6  Hand over and finish','Return change, give invoice and goods. Tap <b>Confirm invoice &amp; hand over goods</b>. Check <b>Sale complete</b>, then <b>Start next sale</b>.')]:
 story.extend([Paragraph(heading,styles['CounterHeading']),Paragraph(body,styles['CounterBody'])])
story.extend([Paragraph('Stop and ask a manager',styles['CounterHeading']),Paragraph('Payment pending or unclear, lost connection, wrong stock or total, return, refund or exchange. <b>Keep the same sale open. Quote the order number. Do not collect again.</b>',styles['CounterBody']),Paragraph('Do not',styles['CounterHeading']),Paragraph('Accept payment screenshots. Hand over while pending. Switch a pending digital sale to cash. Create a replacement after an error. Use a separate QR/payment link. Share accounts or change stock yourself.',styles['CounterBody']),Paragraph('Optional email',styles['CounterHeading']),Paragraph('<b>Submitted to email service</b> does not mean delivered. Use <b>Resend email confirmation</b> if needed. Email does not replace the handwritten invoice.',styles['CounterBody'])])
SimpleDocTemplate(str(OUT/'counter-checklist.pdf'),pagesize=A4,leftMargin=42,rightMargin=42,topMargin=36,bottomMargin=36).build(story)
print('Created DOCX and counter checklist. Render DOCX to create matching PDF and inspect every page.')
